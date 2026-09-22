# -*- coding: utf-8 -*-
"""渲染前保底底片 + 「这一篇该做的都做了吗」提醒 单元/集成测试 (v28.70/28.71, 2026-09-22)

用户侧**真实事故一**: 表格做完了(没做表注/正文), 直接去"出稿"渲染 —— 出来的 PDF 里表格页
还是英文。查依赖后确认: 渲染一路(server/patches/engine/seg_*)从不读表格产物, PDF 的表格页
**本来就保持英文**, 表格译文走附录 DOCX —— 所以"表格没做完就渲染"不会让 PDF 出问题, 只意味着
附录还没生成(随时可重做)。判据于是从"你碰过的做完了吗"(内存清单: 重启即失忆, 也管不了压根
没碰过的那块)换成"**这一篇该做的都做了吗**", 且**只提醒不拦** —— 任意顺序 · 齐了就过 ·
缺了提醒。

用户侧**真实事故二**: 渲染之后原始译稿就散在 TSV 与 #S 块里(机读格式, 人事后想核很费劲),
所以渲染前要留一份**人读存档点**(底片); 底片做不出来时 v28.70 起**直接中止出稿** —— 放它
渲染的后果是 PDF 照样出来且与正常那份无异, 想补底片只能重走一遍, 白多一份 PDF 等人去删。

本套件锁死五件事:

  ① import 不许有副作用 —— `mk_appendix` 末尾原本是裸 `main()`: `mk_ledger` 复用它的
     排版函数时**一 import 就真生成 appendix_tables_zh.docx**, 破坏"只读既有文件"契约。
     断言成对: (a) import 后产物**不在**; (b) 显式调 main() 后产物**在** —— (b) 证明 fixture
     是够的, (a) 才不是空转。
  ② 底片内容 —— 正文逐段「原文/译文」交错; 表格保持**真表格**(可翻格上原文下译文,
     无译文格只留原文); 表注从 notes_zh.json 回填; 落点 <PROJ>\\ledger\\<论文名>_译稿.docx;
     且**只读**既有产物(回包/manifest/notes 字节不变)。
  ③ 底片失败**拦在渲染之前**(v28.70 改判) —— 做不出来就抛 `LedgerFailed`, 出稿中止、报错提醒,
     渲染工序一步都不许走。只在会触发重渲染的任务上做(其它只是排版)。
  ④ 产物判据与篇名归属(纯函数) —— done(产物齐且不比输入旧) / stale(产物比输入旧 = 装配过
     新一轮, 要重做) / todo(产物缺) 全看**产物**, 不存内存(重启面板、隔天回来都不失忆);
     篇名由装配器烙进 manifest(mk_job.py --paper / env P2Z_BODY_NAME, 表注沿用同目录那份),
     与正文那篇一致才算本篇的, 对不上的标 foreign 且**不进提醒**(那是别人的表)。
  ⑤ 提醒而非拦截(HTTP 端到端) —— 缺了照旧出稿 + 台账记「已出稿(缺: X)」+ 日志黄条;
     齐了就过; 非渲染任务(表格/表注)出稿不提醒; 底片失败仍然**硬拦**(v28.70 未被本次改判)。
     逃生门 force / blocked / 本轮清单那条线已整条删除(没有拦截就不需要它)。

不打网络、不动剪贴板、不跑真门禁: run_job 用替身, 面板起在测试进程内的临时端口上。
运行: venv python test_pre_render_ledger.py, 退出码 0=全过
"""
import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = os.path.dirname(HERE)
TABLE_PIPE = os.path.join(TOOLS, "table_pipe")
for p in (TOOLS, TABLE_PIPE):
    if p not in sys.path:
        sys.path.insert(0, p)

# 环境必须在 import 之前落定: 这批模块在**模块层**读 D / inbox / 论文名。
TMP = tempfile.mkdtemp(prefix="p2z_prl_")
TBL = os.path.join(TMP, "tbl")
INBOX = os.path.join(TMP, "inbox")
BODY = "payload_demo"
PDF = os.path.join(TMP, "Demo et al. - 2026 - Test Paper.pdf")
os.makedirs(TBL, exist_ok=True)
os.makedirs(INBOX, exist_ok=True)
for k, v in (("P2Z_PROJ", TMP), ("P2Z_TABLE_DIR", TBL), ("P2Z_INBOX", INBOX),
             ("P2Z_BODY_NAME", BODY), ("P2Z_BODY_PDF", PDF)):
    os.environ[k] = v

import mk_appendix as MA      # noqa: E402  (import 副作用见 §①)
import mk_ledger as ML        # noqa: E402
import mk_notes_job as MN     # noqa: E402  (篇名沿用: _paper_from_job)
import watch_clip as WC       # noqa: E402
import panel as PN            # noqa: E402

APX = os.path.join(TBL, "appendix_tables_zh.docx")
FLAG = os.path.join(TBL, "after_ran.flag")     # 出稿工序"真的跑过"的自证文件(见 §③)
MAN = os.path.join(TBL, "job_manifest.json")
LEDGER = PN.VENDOR_LEDGER
BODY_TEXT = "#S1\n你好世界。\n"


def _w(path, text):
    with io.open(path, "w", encoding="utf-8") as f:
        f.write(text)


def stamp(paper, mt=None):
    """改写 job_manifest.json 里的篇名(v28.71), 并落一个**明确的** mtime ——
    同刻两次写会让"产物比输入旧"的判据抖动, 测试不能靠运气。paper=None 表示没烙。"""
    with io.open(MAN, encoding="utf-8") as f:
        man = json.load(f)
    if paper is None:
        man.pop("paper", None)
    else:
        man["paper"] = paper
    _w(MAN, json.dumps(man, ensure_ascii=False))
    t = time.time() if mt is None else mt
    os.utime(MAN, (t, t))


def fixture():
    """一套最小但**完整**的数据: 表格(有回包) + 表注 + 正文(有回包), 以及 mk_appendix 的
   4 个输入(齐了才证明"import 不生成产物"这条断言不是空转)。"""
    man = {"tables": {"table_10_1": {"header": ["Genera (n)", "Life form"],
                                     "rows": [["10", "Tree"]]}},
           "units": [{"id": "k001", "table": "table_10_1", "r": 0, "c": 0,
                      "orig": "Genera (n)", "ph": {}},
                     {"id": "k002", "table": "table_10_1", "r": 1, "c": 0,
                      "orig": "10", "ph": {}},
                     {"id": "k003", "table": "table_10_1", "r": 1, "c": 1,
                      "orig": "Tree", "ph": {}}]}
    _w(os.path.join(TBL, "job_manifest.json"), json.dumps(man, ensure_ascii=False))
    _w(os.path.join(TBL, "job_response.tsv"), "k001\t属（数量）\nk003\t乔木状\n")
    _w(os.path.join(TBL, "notes_manifest.json"), json.dumps(
        {"units": [{"id": "n001", "table": "table_10_1", "role": "note",
                    "orig": "Table 10.1 ..."}]}, ensure_ascii=False))
    _w(os.path.join(TBL, "notes_zh.json"), json.dumps(
        {"table_10_1": {"note": "表注译文", "foot": "注a译文"}}, ensure_ascii=False))
    for n in ("table_10_1", "table_10_2"):
        _w(os.path.join(TBL, n + ".tsv"), "Genera (n)\tLife form\n10\tTree\n")
        _w(os.path.join(TBL, n + ".zh.tsv"), "属（数量）\t生活型\n10\t乔木状\n")
    _w(os.path.join(INBOX, BODY + ".txt"),
       "[文档] Demo\n[任务] 把下列每个 #S 段落译成简体中文。\n#S1\nHello world.\n#S2\nSecond one.\n")
    _w(os.path.join(INBOX, BODY + ".manifest.json"), json.dumps(
        {"name": BODY, "pdf": PDF,
         "items": [{"key": "S1", "parts": [{"page": 1, "seg": 0}]},
                   {"key": "S2", "parts": [{"page": 1, "seg": 1}]}]}))
    _w(os.path.join(TBL, BODY + "_response.tsv"), BODY_TEXT + "#S2\n第二段。\n")


def http(base, path, body=None):
    req = urllib.request.Request(
        base + path, data=json.dumps(body).encode("utf-8") if body is not None else None,
        headers={"Content-Type": "application/json"} if body is not None else {})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return json.loads(e.read().decode("utf-8"))


def ledger_text():
    with io.open(LEDGER, encoding="utf-8") as f:
        return f.read()


def job_of(pred):
    return next(j for j in WC.available_jobs() if pred(j))


def arm(job, text):
    """把"这次检查"的状态按真实流程摆好(③ 通过后服务器记的就是这三样)。"""
    ids = WC.manifest_ids(job)
    with PN.LOCK:
        PN.STATE["checked"] = (job, ids, {i: "译文" for i in ids})
        PN.STATE["sha"] = PN._sha(text)


def main():                                     # noqa: C901
    passed = failed = 0

    def check(name, cond, detail=""):
        nonlocal passed, failed
        if cond:
            passed += 1
            print(f"  PASS {name}")
        else:
            failed += 1
            print(f"  FAIL {name} {detail}")

    fixture()

    # ---- ① import 无副作用: 一 import 就真生成 docx 是实测踩过的坑 ----
    check("① import mk_appendix 不落 appendix_tables_zh.docx",
          not os.path.exists(APX), APX)
    with contextlib.redirect_stdout(io.StringIO()):
        MA.main()                               # 显式调用才该生成(证明 fixture 够用)
    check("① 显式 main() 才生成(故上一条不是空转)", os.path.exists(APX), APX)
    os.remove(APX)

    # ---- ② 底片内容 + 只读契约 ----
    watched = [os.path.join(TBL, "job_response.tsv"),
               os.path.join(TBL, BODY + "_response.tsv"),
               os.path.join(TBL, "notes_zh.json"),
               os.path.join(TBL, "job_manifest.json")]
    before = {p: io.open(p, "rb").read() for p in watched}
    out = ML.build(log=lambda *_: None)
    want = os.path.join(TMP, "ledger", "Demo et al. - 2026 - Test Paper_译稿.docx")
    check("② 落点 = <PROJ>\\ledger\\<论文名>_译稿.docx", out == want, out)
    check("② 既有产物一个字节都没动",
          all(io.open(p, "rb").read() == before[p] for p in watched))
    real_jobs = WC.JOBS
    WC.JOBS = []                               # 一个 manifest 都没有 = 没有可固化的任务
    check("② 没有任务 manifest -> 不出底片(返回 None)",
          ML.build(log=lambda *_: None) is None)
    WC.JOBS = real_jobs

    from docx import Document
    d = Document(out)
    ps = [p.text for p in d.paragraphs if p.text.strip()]
    check("② 首行标题 = 论文名", ps and ps[0] == "Demo et al. - 2026 - Test Paper", ps[:1])
    check("② 正文逐段「原文/译文」都在", "Hello world." in ps and "你好世界。" in ps, ps)
    check("② 正文按段号升序(原文在前、译文在后)",
          ps.index("Hello world.") < ps.index("你好世界。") < ps.index("Second one."), ps)
    check("② 表注从 notes_zh.json 回填", any(t.startswith("表注: 表注译文") for t in ps), ps)
    check("② 表格是真表格(不是文字块)", len(d.tables) == 1, len(d.tables))
    t = d.tables[0]
    check("② 可翻格: 上原文(灰) / 下译文",
          [p.text for p in t.rows[0].cells[0].paragraphs] == ["Genera (n)", "属（数量）"],
          [p.text for p in t.rows[0].cells[0].paragraphs])
    check("② 无译文格只留原文(不编造)",
          [p.text for p in t.rows[1].cells[0].paragraphs] == ["10"],
          [p.text for p in t.rows[1].cells[0].paragraphs])

    # ---- ③ 底片: 失败要抛(原因带上) / 成功报路径 / 只在渲染任务上做 ----
    real_build = ML.build
    ML.build = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom"))
    logs = []
    try:
        WC._snapshot(log=logs.append)
        check("③ 底片炸了必须抛 LedgerFailed(不许吞)", False, "没抛, 出稿会被放行")
    except WC.LedgerFailed as e:
        check("③ 底片炸了抛 LedgerFailed, 且原因原样带出", "boom" in str(e), str(e))
    ML.build = real_build
    logs = []
    out = WC._snapshot(log=logs.append)
    check("③ 正常时把产物路径报出来", bool(out) and out.endswith("_译稿.docx"), out)

    def fake_job(render):
        return {"label": "正文 x", "resp": "fake.tsv", "blocks": True, "render": render,
                "cmd": lambda: [sys.executable, "-c", "print('PASS ok')"],
                "after": [[sys.executable, "-c",
                           "open(r'%s','a').write('x')" % FLAG]],
                "out": lambda: os.path.join(TBL, "fake_out.docx")}

    # 拦得彻底 = 出稿工序一步没走(用 after 里那条命令写标记文件来自证)
    ML.build = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom"))
    if os.path.exists(FLAG):
        os.remove(FLAG)
    try:
        r = WC.run_job(fake_job(True), ["S1"], {"S1": "甲"}, log=lambda *_: None)
        check("③ 底片失败 -> run_job 抛 LedgerFailed(不返回产物)", False, r)
    except WC.LedgerFailed:
        check("③ 底片失败 -> run_job 抛 LedgerFailed(不返回产物)", True)
    check("③ 底片失败拦得彻底: 出稿工序一步没走(没产物)", not os.path.exists(FLAG))
    ML.build = real_build

    calls = []
    real_snap, WC._snapshot = WC._snapshot, (lambda log=print: calls.append(1))
    r = WC.run_job(fake_job(True), ["S1"], {"S1": "甲"}, log=lambda *_: None)
    check("③ 渲染任务: 门禁过后、工序之前存底片", r and len(calls) == 1, calls)
    calls.clear()
    r = WC.run_job(fake_job(False), ["S1"], {"S1": "甲"}, log=lambda *_: None)
    check("③ 非渲染任务: 不存底片(它只是排版)", r and not calls, calls)
    WC._snapshot = real_snap

    # ---- ④ 产物判据 + 篇名归属(纯函数; 判据在产物上, 不落内存 -> 重启不失忆) ----
    tbl_job = job_of(lambda j: j["label"].startswith("表格正文"))
    note_job = job_of(lambda j: j["label"].startswith("表注"))
    body_job = job_of(lambda j: j.get("render"))
    check("④ 渲染任务靠显式标记认(不靠标签猜)", PN._is_render(body_job)
          and not PN._is_render(tbl_job) and not PN._is_render(note_job))
    check("④ 短名: 正文带论文全名也只显示「正文」", PN.short_of(body_job["label"]) == "正文")
    check("④ 短名: 表格/表注原样", PN.short_of("表格正文") == "表格正文"
          and PN.short_of("表注") == "表注")
    check("④ 判据不靠内存: STATE 里没有「哪块没做」的字段(重启面板不失忆)",
          "round" not in PN.STATE, sorted(PN.STATE))

    check("④ 产物齐且不比输入旧 -> done", WC.released(tbl_job)[0] == "done",
          WC.released(tbl_job))
    check("④ 产物缺 -> todo(正文还没出过稿)", WC.released(body_job)[0] == "todo",
          WC.released(body_job))

    todos = PN.paper_tasks()
    check("④ 本篇待办 = 工作目录里有的每一块",
          [t["short"] for t in todos] == ["表格正文", "表注", "正文"], todos)
    check("④ 状态按产物给(表格/表注 done, 正文 todo)",
          [t["state"] for t in todos] == ["done", "done", "todo"], todos)
    check("④ 没烙篇名 -> 标 unlabeled 但仍算本篇(一个工作目录 = 一篇)",
          all(t["unlabeled"] and not t["foreign"] for t in todos[:2]), todos)
    check("④ 齐了就过: 没有要提醒的", PN.pending_of(todos) == [], PN.pending_of(todos))

    # ① 复制仍落一行外发台账(它是外发历史, 不再是"进清单"的开关)
    if os.path.exists(LEDGER):
        os.remove(LEDGER)
    sent = PN.note_sent(tbl_job, 3)
    check("④ ① 复制落一行外发台账(外发历史)", "复制\t表格正文" in ledger_text(), ledger_text())
    check("④ note_sent 回带本篇待办(供前端刷新)",
          [t["short"] for t in sent] == ["表格正文", "表注", "正文"], sent)

    # 装配过新一轮 = 产物比输入旧 -> stale(不是 done)
    stamp(None, mt=time.time() + 60)
    check("④ 产物比输入旧 -> stale(装配过新一轮, 要重做)",
          WC.released(tbl_job)[0] == "stale", WC.released(tbl_job))
    check("④ stale 的进提醒", [t["short"] for t in PN.pending_of(PN.paper_tasks())]
          == ["表格正文"], PN.pending_of(PN.paper_tasks()))

    # 装配器把篇名烙进 manifest: --paper 优先, 面板 _build 走 env 那条
    def mk_job(*extra):
        p = subprocess.run([sys.executable, os.path.join(TABLE_PIPE, "mk_job.py")] + list(extra),
                           cwd=TBL, capture_output=True, text=True, encoding="utf-8",
                           errors="replace", env=dict(os.environ, P2Z_TABLE_DIR=TBL))
        check("④ 装配器跑得通(退出码 0)", p.returncode == 0, (p.stderr or "")[-300:])
        with io.open(MAN, encoding="utf-8") as f:
            return json.load(f), (p.stdout or "")

    man, so = mk_job("--paper", "论文甲")
    check("④ --paper 烙进 manifest[paper], 且优先于 env", man.get("paper") == "论文甲",
          man.get("paper"))
    check("④ 篇名在装配输出里说一句(人看得见)", "论文甲" in so, so[-200:])
    check("④ 重跑装配 -> 旧产物变 stale", WC.released(tbl_job)[0] == "stale",
          WC.released(tbl_job))
    man, _ = mk_job()
    check("④ 面板 _build 那条路: env P2Z_BODY_NAME 也烙得上", man.get("paper") == BODY,
          man.get("paper"))
    check("④ 表注沿用同目录 job_manifest.json 的篇名(表注手工跑, 不另要参数)",
          MN._paper_from_job() == BODY, MN._paper_from_job())
    stamp(BODY)
    tl = {x["short"]: x for x in PN.paper_tasks()}["表格正文"]
    check("④ 烙了本篇篇名 -> 归属本篇(不标 unlabeled)", not tl["unlabeled"], tl)
    check("④ 篇名缀进显示名(下拉/台账看得出是哪篇的)",
          any(j["label"] == "表格正文 " + BODY for j in WC.available_jobs()),
          [j["label"] for j in WC.available_jobs()])

    # 烙着别篇的表: 仍列出来让人看得见, 但不算成"你欠的"
    stamp("论文乙")
    tl = {x["short"]: x for x in PN.paper_tasks()}["表格正文"]
    check("④ 烙的篇名与正文那篇不一致 -> foreign(别人的表)",
          tl["foreign"] and not tl["unlabeled"], tl)
    check("④ foreign 不进提醒", PN.pending_of(PN.paper_tasks()) == [],
          PN.pending_of(PN.paper_tasks()))
    stamp(None)
    check("④ 没烙篇名 -> 表注也不编造篇名(留空)", MN._paper_from_job() == "",
          MN._paper_from_job())

    # ---- ⑤ HTTP 端到端: 缺了提醒(不拦) / 齐了就过 / 底片失败仍硬拦 ----
    srv = ThreadingHTTPServer(("127.0.0.1", 0), PN.H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = "http://127.0.0.1:%d/" % srv.server_address[1]
    opened = []
    real_run, real_beep = PN.wc.run_job, PN.wc.beep
    had_open, real_open = hasattr(PN.os, "startfile"), getattr(PN.os, "startfile", None)

    def stub_run(job, ids, got, log=print):
        log("STUB 出稿(测试替身)")
        return "OUT"

    PN.wc.run_job, PN.wc.beep = stub_run, (lambda ok=True: None)
    PN.os.startfile = opened.append
    try:
        # (a) 这一篇还有没出稿的 -> 提醒 + 台账, 但**照旧出稿**(不再硬拦)
        stamp(BODY)
        os.remove(os.path.join(TBL, "table_10_1.zh.tsv"))      # 表格正文还没出稿
        with PN.LOCK:
            PN.STATE["checked"] = None
        arm(body_job, BODY_TEXT)
        r = http(base, "api/commit", {"text": BODY_TEXT})
        check("⑤ 缺了照旧出稿(顺序随你, 不再硬拦)", r["ok"] is True and r["out"] == "OUT", r)
        check("⑤ 提醒点名「谁没做」", r.get("warn") == "表格正文", r.get("warn"))
        check("⑤ 提醒写进日志(并说清为什么不要紧)",
              any("没出稿" in l for l in r["log"]) and any("顺序随你" in l for l in r["log"]),
              r["log"])
        check("⑤ 台账记「已出稿(缺: 表格正文)」—— 不是静默放行",
              "已出稿(缺: 表格正文)" in ledger_text(), ledger_text())
        check("⑤ 出稿后清掉已检查状态(防双击重复出稿)", PN.STATE["checked"] is None)
        check("⑤ 待办随产物走: 表格正文现在是 todo",
              [t["state"] for t in r["todos"] if t["short"] == "表格正文"] == ["todo"],
              r["todos"])

        # (b) 非渲染任务出稿不提醒(它只是排版, 不是这一轮的收尾动作)
        arm(tbl_job, "k001\t译文\n")
        r = http(base, "api/commit", {"text": "k001\t译文\n"})
        check("⑤ 非渲染任务出稿: 不提醒(谁先谁后随你)", r["ok"] is True and not r["warn"], r)

        # (c) 齐了就过(两张表的重做产物都补上 —— 缺一张就算没齐)
        for n in ("table_10_1", "table_10_2"):
            _w(os.path.join(TBL, n + ".zh.tsv"), "属（数量）\t生活型\n10\t乔木状\n")
        check("⑤ 补齐产物 -> 没有要提醒的(齐了就过)",
              PN.pending_of(PN.paper_tasks()) == [], PN.pending_of(PN.paper_tasks()))
        arm(body_job, BODY_TEXT)
        r = http(base, "api/commit", {"text": BODY_TEXT})
        check("⑤ 齐了就过: 出稿且无提醒", r["ok"] is True and not r["warn"], r)

        # (d) 底片没做出来: 仍然硬拦(v28.70 的判, 与本次改判无关)
        def stub_ledger_fail(job, ids, got, log=print):
            raise PN.wc.LedgerFailed("底片生成失败(RuntimeError('boom'))")

        arm(body_job, BODY_TEXT)
        opened[:] = []
        PN.wc.run_job = stub_ledger_fail
        r = http(base, "api/commit", {"text": BODY_TEXT})
        check("⑤ 底片没做成 -> 拦下出稿并报错(原因原样提醒)",
              r["ok"] is False and "底片" in r["error"] and "boom" in r["error"], r)
        check("⑤ 底片没做成: 渲染工序一步没走(没开文档)", not opened, opened)
        check("⑤ 底片没做成也记台账(未出稿 + 原因)",
              "未出稿" in ledger_text() and "底片没做成" in ledger_text(), ledger_text())
        check("⑤ 底片没做成: 已检查状态保住(修好后可直接重试 ④, 不用重贴/重跑 ③)",
              PN.STATE["checked"] is not None)
        PN.wc.run_job = stub_run
        r = http(base, "api/commit", {"text": BODY_TEXT})
        check("⑤ 修好后重试 ④ 一次就放行(没有多余的 PDF 要删)",
              r["ok"] is True and r["out"] == "OUT", r)

        # (e) 原有契约: 没先 ③ 检查 -> 拒绝
        with PN.LOCK:
            PN.STATE["checked"] = None
        r = http(base, "api/commit", {"text": "x"})
        check("⑤ 没先 ③ 检查 -> 仍拒绝(原有契约没被绕过)",
              r["ok"] is False and "检查" in r["error"], r)

        # (f) 前端接线: 待办行在; 逃生门那条线整条消失
        pg = PN.PAGE
        for frag in ('id="todoStat"', "renderTodos", "⟳ 要重做",
                     "已出稿, 但这一篇还有没出稿的"):
            check("⑤ 前端接线含 %s" % frag, frag in pg)
        for frag in ('id="forceBtn"', 'id="roundStat"', "renderRound", "force:!!force"):
            check("⑤ 前端已删 %s(没有拦截就不需要逃生门)" % frag, frag not in pg)
    finally:
        PN.wc.run_job, PN.wc.beep = real_run, real_beep
        if had_open:
            PN.os.startfile = real_open
        else:
            del PN.os.startfile
        srv.shutdown()

    shutil.rmtree(TMP, ignore_errors=True)
    print("\n%s: %d passed, %d failed" % (os.path.basename(__file__), passed, failed))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
