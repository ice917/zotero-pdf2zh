# -*- coding: utf-8 -*-
"""附录排版器: table_10_*.zh.tsv + notes_zh.json -> 横版 A4 DOCX(表格中文对照附录)

- 横版 A4: 13 列宽表放不进竖版(原书就是整页侧倒排的表)
- 表头双语: 中文(粗体) + 第二行原书英文列名(小号灰斜体), 便于对照原书
- 首行 tblHeader: 跨页自动重复表头(T2 66 行跨多页)
- 列宽: 按各列最长内容视觉长度(CJK=2) sqrt 加权后归一化到版心, 防长文本列挤死代码列
- 表注/表下脚注排在表下; notes_zh.json 缺失时显红字占位, 不静默

产物: appendix_tables_zh.docx
"""
import csv
import io
import json
import math
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

SD = os.path.dirname(os.path.abspath(__file__))          # 脚本目录(本件所在)
D = os.environ.get("P2Z_TABLE_DIR") or SD                # 工作目录(数据所在), 缺省=脚本目录
TABLES = [("table_10_1", "表 10.1", "Table 10.1"), ("table_10_2", "表 10.2", "Table 10.2")]
GRAY = RGBColor(0x80, 0x80, 0x80)
RED = RGBColor(0xC0, 0x00, 0x00)


def vlen(s):
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)


def style_run(r, size, bold=False, italic=False, color=None):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if color is not None:
        r.font.color.rgb = color


def add_para(doc, text, size=10.5, bold=False, italic=False, color=None,
             align=None, after=6, before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    style_run(p.add_run(text), size, bold, italic, color)
    return p


def set_cell(cell, text, size=9, bold=False, italic=False, color=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    style_run(p.add_run(text), size, bold, italic, color)


def header_cell(cell, zh, en):
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    style_run(p1.add_run(zh), 9, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    style_run(p2.add_run(en), 7.5, italic=True, color=GRAY)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


def fixed_layout(tbl):
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tbl._tbl.tblPr.append(lay)
    m = OxmlElement("w:tblCellMar")
    for side, w in (("left", 40), ("right", 40), ("top", 14), ("bottom", 14)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(w))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tbl._tbl.tblPr.append(m)


def widths_for(hdr_en, rows, usable):
    n = len(hdr_en)
    base = []
    for c in range(n):
        m = 0
        for r in rows:
            m = max(m, min(vlen(r[c] if c < len(r) else ""), 26))
        m = max(m, min(vlen(hdr_en[c]), 20))
        base.append(16 + 9 * math.sqrt(m))
    s = sum(base)
    return [b / s * usable for b in base]


def load_tsv(name):
    with io.open(os.path.join(D, name), encoding="utf-8") as f:
        return [r for r in csv.reader(f, delimiter="\t")]


def main():
    notes = {}
    npath = os.path.join(D, "notes_zh.json")
    if os.path.exists(npath):
        with io.open(npath, encoding="utf-8") as f:
            notes = json.load(f)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.2)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    usable = sec.page_width.pt - sec.left_margin.pt - sec.right_margin.pt

    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(9)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    add_para(doc, "表格中文对照附录", 16, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_para(doc, "Reproductive Biology of Cactaceae · 表 10.1 / 表 10.2", 10.5,
             italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    add_para(doc, "说明: 本附录为原书两表的中文对照。物种学名、命名人、分类代码"
                  "(如 SC、SI、GSI、SHFT、CE)、数值与统计量均按学术惯例保留原文, "
                  "其含义见各表表注; 表头中文下方附原书英文列名, 便于对照。",
             9, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4)

    for ti, (tb, zh_label, en_label) in enumerate(TABLES):
        if ti:
            doc.add_page_break()
        zh = load_tsv(tb + ".zh.tsv")
        en = load_tsv(tb + ".tsv")
        hdr_zh, hdr_en, data = zh[0], en[0], zh[1:]
        add_para(doc, "%s · %s" % (zh_label, en_label), 13, bold=True, before=8, after=6)

        tbl = doc.add_table(rows=1 + len(data), cols=len(hdr_zh))
        tbl.style = "Table Grid"
        fixed_layout(tbl)
        repeat_header(tbl.rows[0])

        widths = widths_for(hdr_en, data, usable)
        for c, w in enumerate(widths):
            tbl.columns[c].width = Pt(w)
            for row in tbl.rows:
                row.cells[c].width = Pt(w)

        for c in range(len(hdr_zh)):
            header_cell(tbl.rows[0].cells[c], hdr_zh[c], hdr_en[c])
        for ri, row in enumerate(data):
            for c in range(len(hdr_zh)):
                set_cell(tbl.rows[ri + 1].cells[c], row[c] if c < len(row) else "")

        nt = notes.get(tb, {})
        if nt.get("note"):
            add_para(doc, "表注: " + nt["note"], 9,
                     align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=8, after=2)
        else:
            add_para(doc, "表注: [待回填 —— 表注译文经 job_notes_doubao.txt 往返后生成]",
                     9, color=RED, before=8, after=2)
        if nt.get("foot"):
            add_para(doc, "注a: " + nt["foot"], 9, before=2)

    out = os.path.join(D, "appendix_tables_zh.docx")
    doc.save(out)

    # 结构自检: 重开数一遍
    chk = Document(out)
    print("附录 -> %s" % out)
    for i, t in enumerate(chk.tables):
        print("  表%d: %d 行 x %d 列" % (i + 1, len(t.rows), len(t.columns)))
        print("  首数据行: %s" % " | ".join(c.text[:12] for c in t.rows[1].cells[:6]))
    print("  表注: %s" % ("已回填" if notes else "占位(红字), 待 notes 往返"))


# 守门: 本件被 import 时(如 mk_ledger.py 复用其排版函数)**不得**执行 main(),
# 否则 import 即真生成 appendix_tables_zh.docx, 破坏调用方的只读契约。
if __name__ == "__main__":
    main()
