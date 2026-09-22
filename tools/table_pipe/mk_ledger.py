# -*- coding: utf-8 -*-
"""渲染前的保底底片: 把人工译稿固化成一份可读 Word (v28.69; v28.70 起失败即中止出稿)

为什么要有它:
  第四步一旦开始渲染, 缓存被注入、产物落盘。若渲染失败、或日后要改字复盘,
  原始译稿散落在 TSV 与 #S 块里 —— 那是机读格式, 人不好核。
  这份底片是**渲染前的人读存档点**: 译稿此刻是什么样, 它就长什么样。

排法依据(业界惯例):
  ① 双语对照 = 两种语言按**自然段落交错**排列 —— 便于校对与后期编辑;
  ② 但该排法"不适合同一表格内编辑" —— 所以表格**保持真表格结构**,
     可翻格内"原文(灰) / 译文"上下排, 不做交错。

两节:
  一、正文         按 #S 升序, 每段 [原文][译文]
  二、表格与表注   每表一个真表格; 表注跟在表后(取 notes_zh.json)

落点: <PROJ>\\ledger\\<论文名>_译稿.docx

契约:
  - **只读**既有文件(回包/manifest/notes), 不写不改任何既有产物;
  - 失败返回 None(无稿可固化)或抛异常 —— 调用方(`watch_clip._snapshot`)**据此中止出稿并
    报错**, 不再放行渲染(v28.70 改判)。原判"副本而已, 失败也不拦"的毛病是: PDF 照样出来
    且与正常那份无异, 想补底片只能重走一遍, 白多一份 PDF 等人去删;
  - 表结构与数字一个字节都不动, 只做排版。

运行: python mk_ledger.py [--paper 论文名]   # 手动补出一份(不动主流程)
"""
import argparse
import csv
import io
import json
import os
import re
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from docx import Document                                  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH              # noqa: E402

import watch_clip as wc                                     # noqa: E402
import mk_appendix as MA                                    # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

LEDGER_DIR = os.path.join(wc.PROJ, "ledger")
_BAD = re.compile(r'[\\/:*?"<>|]')


# ---------------- 读既有产物 ----------------

def tsv_map(path):
    """回包 TSV -> {"k001": 译文}。表格/表注任务都是两列(k id\\t译文)。"""
    out = {}
    if not os.path.exists(path):
        return out
    with io.open(path, encoding="utf-8") as f:
        for row in csv.reader(f, delimiter="\t"):
            if len(row) >= 2 and row[0].strip():
                out[row[0].strip()] = row[1]
    return out


def body_map(path):
    """正文回包 -> {"S1": 译文}。复用 watch_clip 的 #S 块口径(与 write_resp 一致)。"""
    if not os.path.exists(path):
        return {}
    return {"S%d" % n: t for n, t in wc._body_blocks(path).items()}


def notes_map():
    """表注译文(由表注任务的出稿工序生成); 没有就空着, 不报错。"""
    p = os.path.join(wc.D, "notes_zh.json")
    if not os.path.exists(p):
        return {}
    try:
        with io.open(p, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def paper_of(jobs):
    """论文名: 优先正文任务锚定的原文 PDF 的 stem(唯一、能对上源文件)。"""
    for j in jobs:
        pdf = j.get("pdf")
        if pdf:
            return os.path.splitext(os.path.basename(pdf))[0]
    return getattr(wc, "BODY_NAME", "paper")


# ---------------- 排版 ----------------

def _doc(paper):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    MA.add_para(doc, paper, 16, bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    MA.add_para(doc, "译稿底片 · 渲染前存档 · %s" % time.strftime("%Y-%m-%d %H:%M"),
                10.5, italic=True, color=MA.GRAY,
                align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    MA.add_para(doc, "本节为人工译稿的渲染前副本: 正文逐段「原文/译文」对照, "
                     "表格保持原表结构(可翻格内上原文、下译文)。数字、拉丁学名、"
                     "占位符一律未改。",
                9, color=MA.GRAY, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
    return doc


def _cell2(cell, orig, zh, size=8):
    """一格: 原文(灰) 在上, 译文(黑) 在下; 无译文只留原文。"""
    MA.set_cell(cell, orig, size, color=MA.GRAY)
    if zh:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = MA.Pt(0)
        MA.style_run(p.add_run(zh), size + 0.5)


def body_section(doc, jobs, log):
    """一、正文。返回写了多少段。"""
    n = 0
    for j in jobs:
        if not j.get("blocks"):
            continue
        units = wc.manifest_units(j)
        if not units:
            log("    (正文任务 %s 的 manifest/载荷读不到, 跳过)" % j["label"])
            continue
        zh = body_map(os.path.join(wc.D, j["resp"]))
        MA.add_para(doc, "一、正文 · %s（%d 段）" % (j["label"], len(units)),
                    13, bold=True, before=6, after=6)
        miss = 0
        for u in units:
            t = zh.get(u["id"], "")
            MA.add_para(doc, u["orig"], 9.5, color=MA.GRAY, before=4, after=0)
            if t:
                MA.add_para(doc, t, 11, after=2)
            else:
                miss += 1
                MA.add_para(doc, "（本段无译文）", 9, color=MA.RED, after=2)
            n += 1
        if miss:
            log("    ! 正文 %d/%d 段没有译文(底片里标红了)" % (miss, len(units)))
        return n
    return 0


def table_section(doc, jobs, log):
    """二、表格与表注。返回建了几张表。"""
    notes = notes_map()
    n = 0
    for j in jobs:
        if j.get("blocks"):
            continue
        mp = wc.manifest_path(j)
        if not os.path.exists(mp):
            continue
        try:
            with io.open(mp, encoding="utf-8") as f:
                man = json.load(f)
        except (OSError, ValueError) as e:
            log("    ! %s 读不了(%r), 跳过" % (os.path.basename(mp), e))
            continue
        tables = man.get("tables") or {}
        if not tables:
            continue                        # 表注任务的 manifest 没有 tables, 只有 units
        zh = tsv_map(os.path.join(wc.D, j["resp"]))
        if n == 0:
            doc.add_page_break()
            MA.add_para(doc, "二、表格与表注", 13, bold=True, before=6, after=6)

        for tb in tables:
            t = tables[tb]
            hdr, rows = t.get("header") or [], t.get("rows") or []
            if not hdr:
                continue
            at = {}
            for u in man.get("units") or []:
                if u.get("table") != tb:
                    continue
                v = zh.get(u.get("id"))
                if v:
                    at[(u.get("r"), u.get("c"))] = v

            MA.add_para(doc, "%s（%d 行 x %d 列）" % (tb, len(rows), len(hdr)),
                        11.5, bold=True, before=8, after=4)
            tbl = doc.add_table(rows=1 + len(rows), cols=len(hdr))
            tbl.style = "Table Grid"
            MA.fixed_layout(tbl)
            MA.repeat_header(tbl.rows[0])

            for c, h in enumerate(hdr):
                _cell2(tbl.rows[0].cells[c], h, at.get((0, c)))
            for ri, row in enumerate(rows):
                for c in range(len(hdr)):
                    _cell2(tbl.rows[ri + 1].cells[c],
                           row[c] if c < len(row) else "",
                           at.get((ri + 1, c)))

            nt = notes.get(tb) or {}
            if nt.get("note"):
                MA.add_para(doc, "表注: " + nt["note"], 9,
                            align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=6, after=2)
            else:
                MA.add_para(doc, "表注: [未译 —— 表注任务未出稿, 或 notes_zh.json 未生成]",
                            9, color=MA.RED, before=6, after=2)
            if nt.get("foot"):
                MA.add_para(doc, "注a: " + nt["foot"], 9, before=2)
            n += 1
    return n


# ---------------- 入口 ----------------

def build(paper=None, log=print):
    """生成底片 -> 产物路径; 无事可做返回 None。只读既有文件。"""
    jobs = wc.available_jobs()
    if not jobs:
        log("    (没有任何任务 manifest, 跳过底片)")
        return None
    paper = _BAD.sub("_", (paper or paper_of(jobs))).strip() or "paper"

    doc = _doc(paper)
    nb = body_section(doc, jobs, log)
    nt = table_section(doc, jobs, log)
    if not nb and not nt:
        log("    (正文与表格都还没有可固化的译稿, 跳过底片)")
        return None

    if not os.path.isdir(LEDGER_DIR):
        os.makedirs(LEDGER_DIR)
    out = os.path.join(LEDGER_DIR, "%s_译稿.docx" % paper)
    doc.save(out)

    chk = Document(out)                     # 结构自检: 重开数一遍(与 mk_appendix 同风格)
    log("    ✓ 底片 -> %s" % out)
    log("      正文 %d 段 / 表 %d 张" % (nb, len(chk.tables)))
    return out


def main():
    ap = argparse.ArgumentParser(description="渲染前保底底片(只读既有产物)")
    ap.add_argument("--paper", default=None, help="论文名(缺省取正文任务锚定的 PDF 名)")
    a = ap.parse_args()
    out = build(a.paper, log=print)
    print("底片: %s" % out if out else "没能生成(没有可固化的译稿)")
    return 0 if out else 1


if __name__ == "__main__":
    sys.exit(main())
